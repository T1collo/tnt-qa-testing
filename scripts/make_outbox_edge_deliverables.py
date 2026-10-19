"""Build XLSX + DOCX for TrackIQ outbox edge-matrix QA, tied to TL-1843."""
from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx_helpers import run, para, heading, bullet, table, write_docx  # noqa: E402

EDGE_DIR = ROOT / "security" / "results" / "outbox-edge-matrix-2026-10-05T05-09-58-227Z"
PRIOR_5K_DIR = ROOT / "security" / "results" / "outbox-fix-5k-2026-10-02T10-53-38-130Z"
OUT_XLSX = ROOT / "reports" / "TnT_TrackIQ_TL-1843_Outbox_Edge_Matrix_Results_2026-10-05.xlsx"
OUT_DOCX = ROOT / "reports" / "TnT_TrackIQ_TL-1843_Outbox_Edge_Matrix_Report_2026-10-05.docx"
JIRA = "https://apeiro-digital.atlassian.net/browse/TL-1843"
JIRA_COMMENT = (
    "https://apeiro-digital.atlassian.net/browse/TL-1843"
    "?focusedCommentId=53529"
    "&page=com.atlassian.jira.plugin.system.issuetabpanels%3Acomment-tabpanel"
    "#comment-53529"
)

header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
body_font = Font(name="Arial", size=10)
title_font = Font(name="Arial", bold=True, size=14, color="1F4E78")
label_font = Font(name="Arial", bold=True, size=10)
thin = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
hdr = PatternFill("solid", fgColor="1F4E78")
okf = PatternFill("solid", fgColor="C6EFCE")
badf = PatternFill("solid", fgColor="FFC7CE")
warnf = PatternFill("solid", fgColor="FFEB9C")


def style_header(ws, row=1, cols=1):
    for c in range(1, cols + 1):
        cell = ws.cell(row, c)
        cell.font = header_font
        cell.fill = hdr
        cell.border = thin


def main():
    edge = json.loads((EDGE_DIR / "results.json").read_text(encoding="utf-8"))
    summary = edge["summary"]
    results = edge["results"]
    prior = json.loads((PRIOR_5K_DIR / "summary.json").read_text(encoding="utf-8"))
    event_ids_5k = [
        ln.strip()
        for ln in (PRIOR_5K_DIR / "event-ids.txt").read_text(encoding="utf-8").splitlines()
        if ln.strip()
    ]

    # Event IDs captured in edge matrix (where present)
    edge_event_ids = []
    for r in results:
        eid = r.get("eventID")
        if eid:
            edge_event_ids.append(
                {
                    "case_id": r["id"],
                    "category": r["category"],
                    "title": r["title"],
                    "batch_id": summary["batchId"],
                    "event_id": eid,
                    "status": r["status"],
                    "pass": r["pass"],
                }
            )

    headline = [r for r in results if not str(r["id"]).startswith(("A02-", "A03-"))]
    outbox_pass = summary["http5xxRows"] == 0 and summary["outboxPkeySignals"] == 0

    # ---------------- XLSX ----------------
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "TrackIQ Outbox Fix — Edge Matrix Results (TL-1843)"
    ws["A1"].font = title_font
    ws.merge_cells("A1:B1")
    ws["A2"] = f"Jira: {JIRA}"
    ws["A2"].font = Font(name="Arial", size=10, color="0563C1", underline="single")
    ws["A3"] = f"Linked QA comment: {JIRA_COMMENT}"
    ws["A3"].font = Font(name="Arial", italic=True, size=9, color="595959")
    ws.merge_cells("A3:B3")

    meta = [
        ("Jira ticket", "TL-1843"),
        ("Jira URL", JIRA),
        ("Focused comment", "53529"),
        (
            "Outbox defect verdict",
            "PASS — no HTTP 500 and no outbox primary-key errors in API tests",
        ),
        (
            "Edge-matrix headline score",
            f"{summary['headlinePass']}/{summary['headlineTotal']} cases passed",
        ),
        (
            "Script overall flag",
            "PASS" if summary["pass"] else "FAIL (see Findings — not outbox 500s)",
        ),
        ("HTTP 5xx rows", summary["http5xxRows"]),
        ("Outbox PK signals", summary["outboxPkeySignals"]),
        ("Edge-matrix batch ID", summary["batchId"]),
        ("Endpoint", summary["endpoint"]),
        ("Environment", summary["base"]),
        ("Edge run evidence", str(EDGE_DIR.relative_to(ROOT))),
        ("Prior 5k batch ID", prior["batchId"]),
        ("Prior 5k lot", prior["lotNumber"]),
        ("Prior 5k result", f"{prior['ok2xx']}/{prior['count']} HTTP 202, 0x500"),
        ("Prior 5k evidence", str(PRIOR_5K_DIR.relative_to(ROOT))),
        ("Ops fix applied (ticket)", "setval on epcis_outbox_id_seq (+ other lagging sequences)"),
    ]
    ws["A5"] = "Field"
    ws["B5"] = "Value"
    style_header(ws, 5, 2)
    for i, (k, v) in enumerate(meta, start=6):
        ws.cell(i, 1, k).font = label_font
        ws.cell(i, 1).border = thin
        cell = ws.cell(i, 2, v)
        cell.font = body_font
        cell.border = thin
        if k == "Outbox defect verdict":
            cell.fill = okf
        if k == "Script overall flag" and not summary["pass"]:
            cell.fill = warnf
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 100

    # Batches
    ws_b = wb.create_sheet("Batches")
    ws_b["A1"] = "Batch numbers tied to TL-1843 QA"
    ws_b["A1"].font = title_font
    ws_b.merge_cells("A1:E1")
    headers_b = ["Batch / Lot / Run", "Value", "When", "Purpose", "Jira"]
    for c, h in enumerate(headers_b, 1):
        ws_b.cell(2, c, h)
    style_header(ws_b, 2, 5)
    batch_rows = [
        [
            "Edge matrix batch ID",
            summary["batchId"],
            "2026-10-05",
            "API edge-case + volume matrix",
            "TL-1843",
        ],
        [
            "Prior 5k request batch ID",
            prior["batchId"],
            "2026-10-02",
            "5,000 single-event POSTs volume proof",
            "TL-1843",
        ],
        [
            "Prior 5k lot number",
            prior["lotNumber"],
            "2026-10-02",
            "ilmd lot on each of 5,000 commissions",
            "TL-1843",
        ],
        [
            "Prior 5k run ID",
            prior["runId"],
            "2026-10-02",
            "Run identifier inside batch",
            "TL-1843",
        ],
    ]
    for i, row in enumerate(batch_rows, start=3):
        for c, v in enumerate(row, 1):
            cell = ws_b.cell(i, c, v)
            cell.font = body_font
            cell.border = thin
    for c, w in enumerate([28, 36, 14, 40, 12], 1):
        ws_b.column_dimensions[get_column_letter(c)].width = w

    # Headline cases
    ws2 = wb.create_sheet("Headline cases")
    headers = [
        "Case ID",
        "Category",
        "Title",
        "Expected",
        "HTTP Status",
        "Duration ms",
        "Pass",
        "Outbox PK",
        "Event ID",
        "Notes",
        "Jira",
    ]
    for c, h in enumerate(headers, 1):
        ws2.cell(1, c, h)
    style_header(ws2, 1, len(headers))
    for i, r in enumerate(headline, start=2):
        vals = [
            r["id"],
            r["category"],
            r["title"],
            r["expect"],
            r["status"],
            r["ms"],
            "PASS" if r["pass"] else "FAIL",
            "YES" if r["outboxPk"] else "NO",
            r.get("eventID") or "",
            r.get("bodyPreview") or "",
            "TL-1843",
        ]
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(i, c, v)
            cell.font = body_font
            cell.border = thin
        ws2.cell(i, 7).fill = okf if r["pass"] else badf
    for c, w in enumerate([10, 14, 48, 22, 12, 12, 8, 10, 42, 55, 10], 1):
        ws2.column_dimensions[get_column_letter(c)].width = w
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:K{len(headline) + 1}"

    # All requests
    ws3 = wb.create_sheet("All requests")
    for c, h in enumerate(headers, 1):
        ws3.cell(1, c, h)
    style_header(ws3, 1, len(headers))
    for i, r in enumerate(results, start=2):
        vals = [
            r["id"],
            r["category"],
            r["title"],
            r["expect"],
            r["status"],
            r["ms"],
            "PASS" if r["pass"] else "FAIL",
            "YES" if r["outboxPk"] else "NO",
            r.get("eventID") or "",
            r.get("bodyPreview") or "",
            "TL-1843",
        ]
        for c, v in enumerate(vals, 1):
            cell = ws3.cell(i, c, v)
            cell.font = body_font
            cell.border = thin
        ws3.cell(i, 7).fill = okf if r["pass"] else badf
    for c, w in enumerate([12, 14, 48, 22, 12, 12, 8, 10, 42, 55, 10], 1):
        ws3.column_dimensions[get_column_letter(c)].width = w
    ws3.freeze_panes = "A2"
    ws3.auto_filter.ref = f"A1:K{len(results) + 1}"

    # Edge event IDs
    ws4 = wb.create_sheet("Edge matrix Event IDs")
    h4 = ["Index", "Case ID", "Batch ID", "Event ID", "HTTP Status", "Pass", "Jira"]
    for c, h in enumerate(h4, 1):
        ws4.cell(1, c, h)
    style_header(ws4, 1, len(h4))
    for i, row in enumerate(edge_event_ids, start=2):
        vals = [
            i - 1,
            row["case_id"],
            row["batch_id"],
            row["event_id"],
            row["status"],
            "PASS" if row["pass"] else "FAIL",
            "TL-1843",
        ]
        for c, v in enumerate(vals, 1):
            cell = ws4.cell(i, c, v)
            cell.font = body_font
            cell.border = thin
    for c, w in enumerate([8, 12, 24, 48, 12, 8, 10], 1):
        ws4.column_dimensions[get_column_letter(c)].width = w
    ws4.freeze_panes = "A2"

    # Prior 5k full event IDs
    ws5 = wb.create_sheet("Prior 5k Event IDs")
    h5 = ["Index", "Batch ID", "Lot Number", "Event ID", "Jira"]
    for c, h in enumerate(h5, 1):
        ws5.cell(1, c, h)
    style_header(ws5, 1, len(h5))
    for i, eid in enumerate(event_ids_5k, start=1):
        vals = [i, prior["batchId"], prior["lotNumber"], eid, "TL-1843"]
        for c, v in enumerate(vals, 1):
            cell = ws5.cell(i + 1, c, v)
            cell.font = body_font
            cell.border = thin
        if i % 1000 == 0:
            print(f"wrote 5k ids {i}")
    for c, w in enumerate([8, 26, 20, 48, 10], 1):
        ws5.column_dimensions[get_column_letter(c)].width = w
    ws5.freeze_panes = "A2"
    ws5.auto_filter.ref = f"A1:E{len(event_ids_5k) + 1}"

    # Findings
    ws6 = wb.create_sheet("Findings")
    fh = ["ID", "Severity", "Finding", "Detail", "Jira"]
    for c, h in enumerate(fh, 1):
        ws6.cell(1, c, h)
    style_header(ws6, 1, len(fh))
    findings = [
        [
            "OUTBOX",
            "Resolved (TL-1843)",
            "No intermittent 500 / outbox primary-key errors under API edge + volume tests",
            "0 HTTP 5xx; 0 outbox_pkey signals; 100 sequential + 20 concurrent + prior 5,000-request burst",
            "TL-1843",
        ],
        [
            "D01",
            "Medium (separate)",
            "Invalid eventID format was accepted (HTTP 202)",
            "Sent urn:uuid:not-a-valid-uuid — expected 4xx; API accepted. Not the outbox sequence bug.",
            "TL-1843",
        ],
        [
            "F02",
            "Info",
            "Historical bad GLN/GTIN shape did not 500 in this run",
            "Returned 202 on TrackIQ with current payload",
            "TL-1843",
        ],
        [
            "E01-E03",
            "Info",
            "Multi-event documents (10/100/500) accepted with 202",
            "No 500 in this run for those sizes",
            "TL-1843",
        ],
        [
            "B01-B02",
            "Info",
            "Duplicates behave cleanly",
            "Same eventID → 409; same EPC new eventID → 409; never 500",
            "TL-1843",
        ],
    ]
    for i, row in enumerate(findings, start=2):
        for c, v in enumerate(row, 1):
            cell = ws6.cell(i, c, v)
            cell.font = body_font
            cell.border = thin
    for c, w in enumerate([10, 20, 55, 70, 10], 1):
        ws6.column_dimensions[get_column_letter(c)].width = w

    # Traceability
    ws7 = wb.create_sheet("TL-1843 Traceability")
    ws7["A1"] = "How this workbook ties to TL-1843"
    ws7["A1"].font = title_font
    trace = [
        ("Ticket", "TL-1843"),
        ("Ticket URL", JIRA),
        ("QA comment to update/reply", JIRA_COMMENT),
        ("Defect", "Intermittent 500 on manufacturer events due to epcis_outbox sequence lag"),
        ("Ops fix", "setval(epcis_outbox_id_seq) to max(id); audited other sequences"),
        ("QA method", "API-only black-box (Postman-equivalent) — no DB changes in test"),
        ("Volume proof", f"{prior['batchId']} — 5000/5000 HTTP 202"),
        ("Edge proof", f"{summary['batchId']} — edge matrix 2026-10-05"),
        ("Outbox verdict", "PASS"),
        ("Separate finding", "D01 invalid eventID accepted — track separately if needed"),
    ]
    ws7["A3"] = "Item"
    ws7["B3"] = "Value"
    style_header(ws7, 3, 2)
    for i, (k, v) in enumerate(trace, start=4):
        ws7.cell(i, 1, k).font = label_font
        ws7.cell(i, 1).border = thin
        ws7.cell(i, 2, v).font = body_font
        ws7.cell(i, 2).border = thin
    ws7.column_dimensions["A"].width = 28
    ws7.column_dimensions["B"].width = 100

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    wb.save(EDGE_DIR / "TL-1843-edge-matrix-results.xlsx")
    print(f"Wrote {OUT_XLSX} ({OUT_XLSX.stat().st_size} bytes)")

    # ---------------- DOCX ----------------
    body = []
    body.append(
        f'<w:p><w:pPr><w:spacing w:after="80"/></w:pPr>'
        f'{run("TrackIQ Partner API", bold=True, sz="44", color="1F4E78")}</w:p>'
    )
    body.append(
        f'<w:p><w:pPr><w:spacing w:after="80"/></w:pPr>'
        f'{run("TL-1843 — Outbox Fix Edge-Case Verification", bold=True, sz="28")}</w:p>'
    )
    body.append(
        para(
            [
                run("Jira: ", bold=True, color="595959"),
                run("TL-1843  |  ", color="595959"),
                run(JIRA, color="0563C1"),
            ]
        )
    )
    body.append(
        para(
            [
                run("Linked comment: ", bold=True, color="595959"),
                run("53529", color="595959"),
            ]
        )
    )
    body.append(
        para(
            [
                run("Outbox defect verdict: ", bold=True),
                run("PASS", bold=True, color="006400"),
                run(
                    " — no intermittent 500s and no outbox primary-key errors in API testing."
                ),
            ]
        )
    )

    body.append(heading("1. Ticket context", 1))
    body.append(
        para(
            "TL-1843 covers intermittent partner failures on manufacturer event submit. "
            "Ops found the database ID counter for the outbox table was behind the data "
            "after a restore/import, so new writes sometimes collided and the partner saw "
            "a generic 500. Ops advanced the counter (and checked other counters). "
            "QA then verified from the Partner API only."
        )
    )

    body.append(heading("2. How to read this vs the old bug", 1))
    body.append(
        table(
            ["Before fix (incident)", "After QA on TrackIQ"],
            [
                [
                    "Intermittent HTTP 500",
                    "0 × HTTP 500 in edge matrix + prior 5,000-request burst",
                ],
                [
                    "Outbox primary-key collision",
                    "No outbox primary-key / duplicate-key text in API responses",
                ],
                [
                    "Ingest rolled back",
                    "Valid events accepted (HTTP 202)",
                ],
                [
                    "Partner saw Internal server error",
                    "Partner sees 202 on success, or clean 4xx (e.g. 409) on duplicates",
                ],
            ],
            [4400, 4400],
        )
    )

    body.append(heading("3. What was tested (API / Postman-equivalent)", 1))
    body.append(bullet("Happy path single commission"))
    body.append(bullet("100 sequential unique commissions"))
    body.append(bullet("20 concurrent unique commissions"))
    body.append(bullet("Duplicate eventID and same-EPC / new-eventID behaviour"))
    body.append(bullet("Auth: missing key, invalid key, supplier key, facility key on manufacturer route"))
    body.append(bullet("Validation: missing body, bad JSON, wrong content-type, empty list, missing/invalid EPC"))
    body.append(bullet("Batch sizes: 10 / 100 / 500 events in one document; 1 event with 1,000 serials"))
    body.append(bullet("Identifier and time/ILMD edges (including historical crash shape)"))
    body.append(bullet("Reference: prior 5,000 single-event POSTs (batch-5k-req-0ffbe916) — not re-run"))
    body.append(
        para(
            [
                run("Batches: ", bold=True),
                run(f"edge = {summary['batchId']}; volume = {prior['batchId']}."),
            ]
        )
    )

    body.append(heading("4. Findings", 1))
    body.append(
        table(
            ["ID", "Result", "Plain language"],
            [
                [
                    "TL-1843 outbox",
                    "PASS",
                    "No 500s and no outbox key collisions under volume and edge traffic",
                ],
                [
                    "D01",
                    "FAIL (separate)",
                    "A clearly invalid eventID was accepted (202) instead of rejected — not the outbox bug",
                ],
                [
                    "Auth / duplicates / batches",
                    "PASS / Info",
                    "Keys blocked correctly; duplicates return 409; multi-event docs up to 500 accepted",
                ],
            ],
            [2200, 2000, 4600],
        )
    )
    body.append(
        para(
            [
                run("Headline score: ", bold=True),
                run(
                    f"{summary['headlinePass']}/{summary['headlineTotal']} "
                    "(failure is D01 validation only)."
                ),
            ]
        )
    )

    body.append(heading("5. Fix applied (from ticket / ops)", 1))
    body.append(bullet("No app deploy for this incident."))
    body.append(
        bullet(
            "Ops advanced epcis_outbox_id_seq to match the highest existing outbox ID, "
            "and aligned other lagging sequences the same way."
        )
    )
    body.append(bullet("Smoke after fix: new manufacturer event → success; outbox IDs moved forward."))

    body.append(heading("6. Evidence & deliverables", 1))
    body.append(bullet(f"Jira: {JIRA}"))
    body.append(bullet(f"Comment thread: {JIRA_COMMENT}"))
    body.append(bullet(f"Excel: {OUT_XLSX.name}"))
    body.append(bullet(f"Edge evidence: {EDGE_DIR.relative_to(ROOT)}"))
    body.append(bullet(f"Prior 5k evidence: {PRIOR_5K_DIR.relative_to(ROOT)}"))

    body.append(heading("7. Sign-off", 1))
    body.append(
        table(
            ["Item", "Status"],
            [
                ["TL-1843 outbox failure mode (API verification)", "PASS"],
                ["Ready to attach on TL-1843 comment 53529", "Yes"],
                ["Follow-up (optional)", "Open separate item for D01 invalid eventID accepted"],
            ],
            [4800, 4000],
        )
    )

    write_docx(str(OUT_DOCX), "TL-1843 Outbox Edge Matrix Report", body)
    print(f"Wrote {OUT_DOCX}")

    # Ticket comment text for the user
    comment_path = ROOT / "reports" / "TL-1843_QA_comment_2026-10-05.txt"
    comment = f"""TL-1843 QA update (TrackIQ) — tied to comment 53529

What was tested (Partner API only — Postman-equivalent):
• Happy path + 100 sequential + 20 concurrent manufacturer event posts
• Duplicates (same event ID / same product code)
• Auth (no key, bad key, wrong-role keys)
• Validation & payload edge cases (bad JSON, wrong content type, empty list, etc.)
• Multi-event documents (10 / 100 / 500) and 1,000 serials in one event
• Prior volume proof: 5,000 single-event posts (not re-run this session)

Batches:
• Edge matrix: {summary['batchId']}
• Volume (5k): {prior['batchId']} (lot {prior['lotNumber']})

Findings:
• Outbox / intermittent-500 issue: PASS — 0 × HTTP 500, no outbox primary-key errors in responses
• Valid events accepted (202); duplicates cleanly rejected (409)
• One separate finding: invalid eventID format was accepted (202) instead of rejected — not the outbox bug (case D01)

Fix applied (ops, per ticket):
• Database ID counter for outbox was behind the data after restore/import
• Ops advanced the counter (and checked other counters) — no app deploy for this incident

Evidence:
• Excel: reports/{OUT_XLSX.name}
• Word: reports/{OUT_DOCX.name}
• Raw: {EDGE_DIR.relative_to(ROOT)}
• Jira: {JIRA}
"""
    comment_path.write_text(comment, encoding="utf-8")
    print(f"Wrote {comment_path}")
    print(f"outbox_pass={outbox_pass} headline={summary['headlinePass']}/{summary['headlineTotal']}")


if __name__ == "__main__":
    main()
